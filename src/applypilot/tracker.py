"""Excel application tracker export.

Generates an .xlsx tracker of every job ApplyPilot has actually applied to,
in the same spirit as the classic LAMP-List-style job-search tracker
(Position / Company / Role / Location / Date Applied / Résumé / Link /
Status / Notes / ...): ApplyPilot fills in what it knows -- what you applied
to, where, when, with which résumé, and (if `applypilot gmail sync` is set
up -- see gmail_status.py) Last Contact and a first guess at Status from
your inbox. Purely personal columns (Notes, Connections?, Latest word,
contact 1, SHADE) are always left for you to fill in by hand.

This is its own file (~/.applypilot/tracker.xlsx), not something that
touches any tracker you already keep by hand -- point a copy/paste or a
formula at it if you want to fold it into an existing sheet.

Regenerating is an upsert, not an overwrite: matched by the Link column, so
re-running `export()` after every new application refreshes the columns
ApplyPilot owns (see `_OWNED` below) without ever touching a cell you've
filled in yourself.
"""

import logging
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from applypilot import config
from applypilot.database import get_connection

logger = logging.getLogger(__name__)

TRACKER_PATH = config.APP_DIR / "tracker.xlsx"
SHEET_NAME = "Applications"

HEADERS = [
    "Position", "Company", "Industry", "Role", "Location", "Location (State)",
    "Date Posted", "Date Applied", "Last Contact", "Connections?", "Cover Letter",
    "Résumé upload?", "Résumé Form?", "Salary Range", "per", "Anticipated Takeaway",
    "Link", "Notes", "Status", "Latest word", "contact 1", "SHADE",
]

# Columns ApplyPilot fills and refreshes on every export -- Last Contact is
# here too since it's an objective Gmail timestamp (see gmail_status.py),
# not a judgment call. Status is handled separately below: it's seeded
# from Gmail but a value you type into the sheet by hand always wins over
# the automatic guess from then on (see the manual-override check in
# export()). Everything else (Notes, Connections?, Latest word, contact 1,
# SHADE) is entirely yours -- ApplyPilot never writes there at all.
_OWNED = {
    "Position", "Company", "Role", "Location", "Location (State)",
    "Date Posted", "Date Applied", "Last Contact", "Cover Letter",
    "Résumé upload?", "Résumé Form?", "Salary Range", "per",
    "Anticipated Takeaway", "Link",
}

_HEADER_FILL = PatternFill("solid", fgColor="B6D7A8")
_HEADER_FONT = Font(name="Calibri", size=12, bold=True)
_LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
_DATE_FMT = 'mmmm" "d", "yyyy'
_CURRENCY_FMT = '"$"#,##0'

_STATUS_OPTIONS = (
    "- -,draft,sent,pending,phone call,interview round 1,interview round 2,"
    "rejection (application),rejection (interview),ignored (application),"
    "ghosted (interview),turned down,offer"
)

_COLUMN_WIDTHS = {
    "Position": 40, "Company": 18, "Industry": 16, "Role": 10,
    "Location": 20, "Location (State)": 14, "Date Posted": 16,
    "Date Applied": 16, "Cover Letter": 22, "Résumé upload?": 30,
    "Résumé Form?": 12, "Salary Range": 12, "per": 8,
    "Anticipated Takeaway": 16, "Link": 28, "Notes": 30, "Status": 22,
}

_ATS_DOMAINS = (
    "myworkdayjobs.com", "greenhouse.io", "icims.com", "lever.co",
    "taleo.net", "bamboohr.com", "smartrecruiters.com", "jobs.gem.com",
    "ashbyhq.com", "jobvite.com", "successfactors.com", "workable.com",
)


def _col(name: str) -> int:
    return HEADERS.index(name) + 1


def _infer_role(title: str | None) -> str:
    return "Intern" if title and "intern" in title.lower() else ""


def _infer_state(location: str | None) -> str:
    if not location:
        return ""
    parts = [p.strip() for p in location.split(",")]
    return parts[-1] if len(parts) > 1 else ""


def _infer_ats(url: str | None) -> str:
    url = (url or "").lower()
    return "ATS" if any(d in url for d in _ATS_DOMAINS) else "no ATS"


def _parse_salary(raw: str | None) -> tuple[float | None, str]:
    """Best-effort: pull one clean number and a per-period guess out of
    free-text salary. Returns (None, "") whenever it can't confidently
    parse a single number -- the raw text is used as-is in that case."""
    if not raw:
        return None, ""
    numbers = re.findall(r"[\d,]+(?:\.\d+)?", raw)
    if len(numbers) != 1:
        return None, ""
    try:
        value = float(numbers[0].replace(",", ""))
    except ValueError:
        return None, ""

    low = raw.lower()
    if "hr" in low or "hour" in low:
        period = "hour"
    elif "yr" in low or "year" in low or "annual" in low:
        period = "year"
    elif "month" in low:
        period = "month"
    else:
        period = ""
    return value, period


def _parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value).replace(tzinfo=None)
    except ValueError:
        return None


def _takeaway_formula(row: int) -> str:
    """Estimated take-home for a ~16-week internship (640 hrs), mirroring
    the classic LAMP-tracker formula: hourly*640, monthly*4, else annual
    pro-rated. Wrapped in IFERROR since Salary Range is sometimes free text."""
    n_ref = f"{get_column_letter(_col('Salary Range'))}{row}"
    o_ref = f"{get_column_letter(_col('per'))}{row}"
    return (
        f'=IFERROR(IF(OR(ISBLANK({n_ref}),ISBLANK({o_ref})),"",'
        f'IF({o_ref}="hour",{n_ref}*640,IF({o_ref}="month",{n_ref}*4,{n_ref}/12/160*640))),"")'
    )


def _row_values(job: dict) -> dict:
    url = job.get("application_url") or job["url"]
    salary_num, per = _parse_salary(job.get("salary"))

    resume_path = job.get("tailored_resume_path")
    resume_name = Path(resume_path).with_suffix(".pdf").name if resume_path else ""

    cl_path = job.get("cover_letter_path")
    cover_letter = Path(cl_path).with_suffix(".pdf").name if cl_path else "N/A"

    return {
        "Position": job.get("title") or "",
        "Company": job.get("site") or "",
        "Role": _infer_role(job.get("title")),
        "Location": job.get("location") or "",
        "Location (State)": _infer_state(job.get("location")),
        "Date Posted": _parse_iso(job.get("discovered_at")),
        "Date Applied": _parse_iso(job.get("applied_at")),
        "Last Contact": _parse_iso(job.get("gmail_last_contact_at")),
        "Cover Letter": cover_letter,
        "Résumé upload?": resume_name,
        "Résumé Form?": _infer_ats(url),
        "Salary Range": salary_num if salary_num is not None else (job.get("salary") or ""),
        "per": per,
        "Link": url,
    }


def _new_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for c, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="left")
    ws.freeze_panes = "A2"
    for name, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(_col(name))].width = width
    return wb, ws


def _refresh_status_validation(ws) -> None:
    for dv in list(ws.data_validations.dataValidation):
        ws.data_validations.dataValidation.remove(dv)
    dv = DataValidation(type="list", formula1=f'"{_STATUS_OPTIONS}"', allow_blank=True)
    status_col = get_column_letter(_col("Status"))
    last_row = max(ws.max_row, 2)
    dv.add(f"{status_col}2:{status_col}{last_row}")
    ws.add_data_validation(dv)


def export(path: Path | None = None) -> Path:
    """(Re)generate the tracker from every job marked applied.

    Upserts by the Link column: ApplyPilot-owned fields (including Last
    Contact, from Gmail) are refreshed on existing rows and new applications
    get appended. Status is seeded from Gmail too, but the moment you type
    a different value into that cell by hand, it's yours -- future exports
    won't touch it again. Notes/Connections?/Latest word/contact 1/SHADE are
    never written by ApplyPilot at all.
    """
    path = path or TRACKER_PATH
    conn = get_connection()
    jobs = [dict(r) for r in conn.execute(
        "SELECT * FROM jobs WHERE apply_status = 'applied' ORDER BY applied_at"
    ).fetchall()]

    ws = None
    if path.exists():
        try:
            wb = openpyxl.load_workbook(path)
            if SHEET_NAME in wb.sheetnames:
                ws = wb[SHEET_NAME]
        except Exception:
            logger.exception("tracker: couldn't open existing %s, recreating", path)

    existing_by_link: dict[str, int] = {}
    if ws is not None:
        link_col = _col("Link")
        for r in range(2, ws.max_row + 1):
            link = ws.cell(row=r, column=link_col).value
            if link:
                existing_by_link[link] = r
    else:
        wb, ws = _new_workbook()

    next_row = ws.max_row + 1 if ws.max_row >= 2 else 2

    for job in jobs:
        url = job.get("application_url") or job["url"]
        values = _row_values(job)
        row = existing_by_link.get(url)
        is_new = row is None
        if is_new:
            row = next_row
            next_row += 1

        for name, value in values.items():
            cell = ws.cell(row=row, column=_col(name))
            cell.value = value
            if name in ("Date Posted", "Date Applied", "Last Contact") and value is not None:
                cell.number_format = _DATE_FMT
            elif name == "Salary Range" and isinstance(value, (int, float)):
                cell.number_format = _CURRENCY_FMT
            elif name == "Link" and value:
                cell.hyperlink = value
                cell.font = _LINK_FONT

        takeaway_cell = ws.cell(row=row, column=_col("Anticipated Takeaway"))
        takeaway_cell.value = _takeaway_formula(row)
        takeaway_cell.number_format = _CURRENCY_FMT

        # Status: seeded/refreshed from Gmail, but a value typed into the
        # sheet by hand permanently overrides the automatic guess -- we only
        # ever write here if the current cell still matches what ApplyPilot
        # itself wrote last time (or is a brand new / still-blank row).
        status_cell = ws.cell(row=row, column=_col("Status"))
        current_status = status_cell.value
        prior_written = job.get("status_last_written")
        gmail_status = job.get("gmail_status")

        if is_new:
            new_value = gmail_status or "- -"
            status_cell.value = new_value
            new_written = new_value
        elif current_status in (None, "", "- -") or current_status == prior_written:
            if gmail_status and gmail_status != current_status:
                status_cell.value = gmail_status
            new_written = gmail_status or current_status
        else:
            new_written = prior_written  # user has taken over this cell -- leave it alone

        if new_written != prior_written:
            conn.execute(
                "UPDATE jobs SET status_last_written = ? WHERE url = ?",
                (new_written, job["url"]),
            )

    conn.commit()
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"
    _refresh_status_validation(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("tracker: exported %d applied job(s) to %s", len(jobs), path)
    return path
